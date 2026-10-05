"""
Shared fixtures. Tests build their own sample files rather than relying
on files you happen to have lying around -- that's what makes this
suite runnable on any machine, including CI, with no setup.
"""

import sys
from pathlib import Path

import pytest
import pymupdf
import docx

# Make `from app...` imports work when pytest is run from backend/
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))


@pytest.fixture(autouse=True)
def _isolated_run_history(tmp_path, monkeypatch):
    """
    Every run_pipeline call appends to the run history. Without this, every
    test run -- unit tests included -- wrote into the REAL
    evaluation_log.jsonl that the UI's "Run history" panel reads, burying
    real runs under hundreds of "Discuss the topic" entries.
    """
    from app import evaluation
    monkeypatch.setattr(evaluation, "EVALUATION_LOG_PATH", str(tmp_path / "evaluation_log.jsonl"))


@pytest.fixture(autouse=True)
def _isolated_figure_cache(tmp_path, monkeypatch):
    """Same idea for figure descriptions: tests must never read descriptions
    cached by real runs, or write fake ones where real runs would find them."""
    from app import doc_parser
    monkeypatch.setattr(doc_parser, "FIGURE_CACHE_PATH", str(tmp_path / "figure_cache.json"))


@pytest.fixture(autouse=True)
def _isolated_auth_db(tmp_path, monkeypatch):
    """Accounts, job ownership and queued uploads created by tests go to a
    throwaway location, never the real researchlens.db / job_uploads."""
    from app import config
    monkeypatch.setattr(config, "AUTH_DB_PATH", str(tmp_path / "researchlens.db"))
    monkeypatch.setattr(config, "UPLOAD_DIR", str(tmp_path / "job_uploads"))


@pytest.fixture(autouse=True)
def _fresh_groq_token_budget():
    """Rate-limit state learned from one test's fake responses must not
    make the next test wait."""
    from app import groq_client
    groq_client._token_budget.clear()


@pytest.fixture
def sample_pdf(tmp_path) -> str:
    """A 2-page PDF with distinct, known text on each page."""
    doc = pymupdf.open()
    page1 = doc.new_page()
    page1.insert_text((72, 72), "Page one discusses machine learning in healthcare.")
    page2 = doc.new_page()
    page2.insert_text((72, 72), "Page two discusses data privacy concerns in medical AI.")
    path = tmp_path / "sample.pdf"
    doc.save(str(path))
    doc.close()
    return str(path)


@pytest.fixture
def sample_docx(tmp_path) -> str:
    """A DOCX with two headed sections, for testing section-based splitting."""
    d = docx.Document()
    d.add_heading("Introduction", level=1)
    d.add_paragraph("This introduces the topic of machine learning in medicine.")
    d.add_heading("Methodology", level=1)
    d.add_paragraph("This describes the research methodology used in the study.")
    path = tmp_path / "sample.docx"
    d.save(str(path))
    return str(path)


@pytest.fixture
def sample_txt(tmp_path) -> str:
    path = tmp_path / "sample.txt"
    path.write_text("A short plain-text research note about neural networks.")
    return str(path)


_PARAGRAPH = (
    "This paragraph is ordinary body text of the paper, long enough to count as a "
    "paragraph rather than a figure label, describing how the evaluation was set up "
    "across five public benchmark datasets and three model sizes in total."
)


@pytest.fixture
def figure_pdf(tmp_path) -> str:
    """
    Two pages, built to cover what figure detection has to tell apart:
      page 1 -- a paragraph, then a VECTOR bar chart (drawn rectangles plus
                axis labels, no embedded image) captioned "Figure 1: ...".
      page 2 -- a sentence that merely refers to Figure 1 (not a caption), a
                RASTER image captioned "Fig. 2 | ...", and a "Figure 3."
                caption with nothing graphical above it (a false match).
    """
    doc = pymupdf.open()

    page1 = doc.new_page()
    page1.insert_textbox(pymupdf.Rect(72, 60, 540, 140), _PARAGRAPH, fontsize=10)
    for i, height in enumerate([60, 100, 80, 130, 40]):
        page1.draw_rect(pymupdf.Rect(110 + i * 60, 330 - height, 150 + i * 60, 330),
                        color=(0, 0, 0), fill=(0.2, 0.4, 0.8))
    page1.draw_line((100, 330), (420, 330))
    page1.insert_text((95, 345), "D1   D2   D3   D4   D5", fontsize=8)
    page1.insert_text((60, 260), "Acc", fontsize=8)
    page1.insert_textbox(pymupdf.Rect(72, 360, 540, 400),
                         "Figure 1: Accuracy of the model across five datasets.", fontsize=10)

    page2 = doc.new_page()
    page2.insert_textbox(pymupdf.Rect(72, 60, 540, 140),
                         "Figure 1 shows that accuracy varies widely. " + _PARAGRAPH, fontsize=10)
    pixmap = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 120, 80), False)
    pixmap.clear_with(180)
    page2.insert_image(pymupdf.Rect(150, 170, 390, 330), pixmap=pixmap)
    page2.insert_textbox(pymupdf.Rect(72, 340, 540, 380),
                         "Fig. 2 | Example input image from the dataset.", fontsize=10)
    page2.insert_textbox(pymupdf.Rect(72, 420, 540, 500), _PARAGRAPH, fontsize=10)
    page2.insert_textbox(pymupdf.Rect(72, 505, 540, 540),
                         "Figure 3. A caption with no figure above it.", fontsize=10)

    path = tmp_path / "figures.pdf"
    doc.save(str(path))
    doc.close()
    return str(path)


@pytest.fixture
def empty_pdf(tmp_path) -> str:
    """A PDF with a page but no text at all -- simulates a scanned/image-only PDF."""
    doc = pymupdf.open()
    doc.new_page()  # blank page, no insert_text call
    path = tmp_path / "empty.pdf"
    doc.save(str(path))
    doc.close()
    return str(path)
